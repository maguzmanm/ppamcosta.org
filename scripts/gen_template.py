import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import json

wb = openpyxl.Workbook()
hf = Font(bold=True, color="FFFFFF", size=11)
hfill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
nf = Font(italic=True, color="888888", size=9)

# ─── Hoja Publicadores ───
ws = wb.active
ws.title = "Publicadores"
headers = [
    "firstName",
    "lastName",
    "marriedLastName",
    "gender",
    "email",
    "phone",
    "maritalStatus",
    "spouseEmail",
    "congregationId",
    "designations",
    "notes",
]
examples = [
    "Maria",
    "Gonzalez",
    "Rojas",
    "F",
    "maria@email.com",
    "+56 9 1234 5678",
    "CASADO",
    "juan@email.com",
    "",
    '["PRECURSOR_REGULAR"]',
    "",
]
notes = [
    "Nombre",
    "Apellido",
    "Apellido casada",
    "M o F",
    "Email",
    "Telefono",
    "SOLTERO|CASADO|DIVORCIADO|SEPARADO|VIUDO",
    "Email conyuge PPAM",
    "ID congregacion (ver hoja Congregaciones)",
    "Designaciones JSON",
    "Notas",
]
widths = [14, 14, 16, 8, 24, 18, 14, 24, 40, 26, 26]

for col, (h, ex, n) in enumerate(zip(headers, examples, notes), 1):
    c = ws.cell(row=1, column=col, value=h)
    c.font = hf
    c.fill = hfill
    c.alignment = Alignment(horizontal="center")
    ws.cell(row=2, column=col, value=ex)
    ws.cell(row=3, column=col, value=n).font = nf
    ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]

# ─── Hoja Ayuda ───
ws2 = wb.create_sheet("Ayuda")
ws2.column_dimensions["A"].width = 40
ws2.column_dimensions["B"].width = 55
help_data = [
    ("CAMPO", "DESCRIPCION"),
    ("firstName", "Nombre del publicador. Obligatorio."),
    ("lastName", "Apellido del publicador. Obligatorio."),
    ("marriedLastName", "Apellido de casada. Solo para mujeres casadas."),
    ("gender", "M = Masculino, F = Femenino."),
    ("email", "Email para inicio de sesion. Vacio = no se crea usuario."),
    ("phone", "Telefono en formato libre."),
    ("maritalStatus", "SOLTERO, CASADO, DIVORCIADO, SEPARADO, VIUDO."),
    ("spouseEmail", "Email del conyuge si ya esta en PPAM. Vacio si no."),
    ("congregationId", "ID de la congregacion. Ver hoja Congregaciones."),
    ("designations", 'JSON: ["ANCIANO","PRECURSOR_REGULAR","SIERVO_MINISTERIAL",...]'),
    ("notes", "Notas adicionales (opcional)."),
]
for i, (k, v) in enumerate(help_data, 1):
    c1 = ws2.cell(row=i, column=1, value=k)
    c2 = ws2.cell(row=i, column=2, value=v)
    if i == 1:
        c1.font = hf
        c1.fill = hfill
        c2.font = hf
        c2.fill = hfill

# ─── Hoja Congregaciones ───
congregations = [
    ("1f63a7fc-6940-4999-9bf5-728d502620f0", "Alto Mirador - San Antonio"),
    ("b3ef5456-4122-41a5-9486-ab55d5ee0306", "Barrancas - San Antonio"),
    ("b9b15417-66ad-43c7-b306-c51f1332572b", "Barrio Oriente - Melipilla"),
    ("49cb3f18-f908-493d-89eb-677fb21cb637", "Brisas de Culipran - Melipilla"),
    ("64ec34cf-5292-4bf0-ae6b-f9f4b0feda94", "Centenario - San Antonio"),
    ("4927859c-5c7b-4168-b516-db90a9364bfe", "Central - Melipilla"),
    ("c54e0fae-d37b-4240-8589-1ac18cd43ae5", "Costa Azul - Cartagena"),
    ("3e758822-2970-4dff-816c-cf8e885881e6", "Costanera - El Quisco"),
    ("029db583-9714-40f9-8e2d-4898f735372a", "Criollo Haitiano Melipilla"),
    ("636ecf00-fb62-417d-b7ba-1624785abcea", "Criollo Haitiano San Antonio"),
    ("c034fade-08e7-417d-b35d-47149cdb686e", "El Canelo - Algarrobo"),
    ("cad483fd-00f1-49cf-905b-24b2754ca75a", "El Tabo"),
    ("88f76fe2-8382-48fa-be53-9d5b7d7edf25", "El Totoral - El Quisco"),
    ("63f475ad-a6cb-48a4-bdf5-34ad86170f37", "El Tranque - Llolleo"),
    ("05af16f6-e994-4543-80b8-e824bbc90963", "Isla Negra"),
    ("6d939a38-7654-411f-9501-6578dbc91ad9", "Litoral - Cartagena"),
    ("38db7727-4738-487e-a804-243e66c41b45", "Lo Abarca - Cartagena"),
    ("186a1977-9cb2-4e5c-860d-3515b139cd2e", "Lomas - Llolleo"),
    ("10e45fca-0c7e-4ad6-944e-eb37606606e4", "Los Boldos - Bollenar"),
    ("258ad838-9d8b-4d5b-bef8-70861140e7ec", "Mirasol - Algarrobo"),
    ("44ff763e-a1e2-4e88-be2e-41817eea1c7c", "Pabellon - Melipilla"),
    ("bcd9e619-f7ac-4555-8bc6-b0fbbdde115d", "Pomaire"),
    ("104b70da-463c-4343-b79b-79510f75b061", "Riberas del Maipo - Santo Domingo"),
    ("095fcf42-e128-49f1-9b2a-a618204918e2", "Sol Poniente - Melipilla"),
    ("b4380cad-4ec6-47ae-b7ad-52867e824c7f", "Valle de Mallarauco - Bollenar"),
    ("32c7c1a1-cc9a-4b90-b153-7a27164c5493", "Villa Alhue"),
]

ws3 = wb.create_sheet("Congregaciones")
ws3.column_dimensions["A"].width = 40
ws3.column_dimensions["B"].width = 40
for col, t in enumerate(["ID", "Nombre"], 1):
    c = ws3.cell(row=1, column=col, value=t)
    c.font = hf
    c.fill = hfill
for i, (cid, name) in enumerate(congregations, 2):
    ws3.cell(row=i, column=1, value=cid)
    ws3.cell(row=i, column=2, value=name)

wb.save("plantilla_importar_publicadores_v2.xlsx")
print(
    f"OK - {len(congregations)} congregaciones en plantilla_importar_publicadores_v2.xlsx"
)

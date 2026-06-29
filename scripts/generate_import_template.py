import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import subprocess

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Publicadores"

header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
note_font = Font(italic=True, color="888888", size=9)

headers = [
    "firstName",
    "lastName",
    "marriedLastName",
    "gender",
    "email",
    "phone",
    "maritalStatus",
    "spouseEmail",
    "congregationId",
    "designations",
    "notes",
]
examples = [
    "María",
    "González",
    "Rojas",
    "F",
    "maria@email.com",
    "+56 9 1234 5678",
    "CASADO",
    "juan@email.com",
    "",
    '["PRECURSOR_REGULAR"]',
    "",
]
notes = [
    "Nombre",
    "Apellido",
    "Apellido casada (opcional)",
    "M o F",
    "Email (opcional)",
    "Teléfono",
    "SOLTERO|CASADO|DIVORCIADO|SEPARADO|VIUDO",
    "Email del cónyuge en PPAM (opcional)",
    "ID de la congregación (ver hoja Congregaciones)",
    "Designaciones JSON (opcional)",
    "Notas (opcional)",
]

for col, (h, ex, n) in enumerate(zip(headers, examples, notes), 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")
    ws.cell(row=2, column=col, value=ex)
    ws.cell(row=3, column=col, value=n).font = note_font

widths = [14, 14, 16, 8, 24, 18, 14, 24, 40, 26, 26]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# ─── Hoja Ayuda ───
ws2 = wb.create_sheet("Ayuda")
ws2.column_dimensions["A"].width = 40
ws2.column_dimensions["B"].width = 55
help_data = [
    ("CAMPO", "DESCRIPCIÓN"),
    ("firstName", "Nombre del publicador. Obligatorio."),
    ("lastName", "Apellido del publicador. Obligatorio."),
    ("marriedLastName", "Apellido de casada. Solo para mujeres casadas."),
    ("gender", "M = Masculino, F = Femenino."),
    ("email", "Email para inicio de sesión. Si se deja vacío no se crea usuario."),
    ("phone", "Teléfono en formato libre."),
    ("maritalStatus", "SOLTERO, CASADO, DIVORCIADO, SEPARADO, o VIUDO."),
    ("spouseEmail", "Email del cónyuge si ya está en PPAM. Si no, dejar vacío."),
    (
        "congregationId",
        "ID de la congregación. Debe existir (ver hoja Congregaciones).",
    ),
    ("designations", 'JSON: ["ANCIANO","PRECURSOR_REGULAR","SIERVO_MINISTERIAL",...]'),
    ("notes", "Notas adicionales (opcional)."),
]
for i, (k, v) in enumerate(help_data, 1):
    c1 = ws2.cell(row=i, column=1, value=k)
    c2 = ws2.cell(row=i, column=2, value=v)
    if i == 1:
        c1.font = header_font
        c1.fill = header_fill
        c2.font = header_font
        c2.fill = header_fill


# ─── Hoja Congregaciones (datos reales desde MySQL) ───
def fetch_congregations():
    try:
        result = subprocess.run(
            [
                "mysql",
                "-u",
                "ppam_user",
                "-pPpamCosta2026",
                "ppamcosta_db",
                "-e",
                "SELECT id, name FROM congregations ORDER BY name",
            ],
            capture_output=True,
            text=True,
            timeout=10,
        )
        lines = result.stdout.strip().split("\n")[1:]
        return [l.split("\t") for l in lines if l.strip()]
    except Exception as e:
        print(f"No se pudo conectar a MySQL: {e}")
        return []


congs = fetch_congregations()
if congs:
    ws3 = wb.create_sheet("Congregaciones")
    ws3.column_dimensions["A"].width = 40
    ws3.column_dimensions["B"].width = 40
    for col, title in enumerate(["ID", "Nombre"], 1):
        c = ws3.cell(row=1, column=col, value=title)
        c.font = header_font
        c.fill = header_fill
    for i, row in enumerate(congs, 2):
        for j, val in enumerate(row[:2]):
            ws3.cell(row=i, column=j + 1, value=val)
    print(f"Hoja 'Congregaciones' agregada con {len(congs)} congregaciones.")

wb.save("plantilla_importar_publicadores.xlsx")
print("Plantilla generada: plantilla_importar_publicadores.xlsx")
